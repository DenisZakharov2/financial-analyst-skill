"""Генерация отчётов: DOCX, PDF, XLSX, PPTX, RTF и проверка текста.

Отчёт содержит только факты из metrics.json: ключевые показатели, таблицы,
допущения, ограничения и контроль качества. Интерпретацию и рекомендации
пишет агент по шаблону templates/report.md.tmpl, а не этот модуль.
"""

from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path
from xml.sax.saxutils import escape

import numpy as np
import pandas as pd

from .core import QualityReport

PAGE_WIDTH_INCHES = 6.5


@dataclass
class ReportData:
    title: str
    subtitle: str = ""
    headline: list[dict] = field(default_factory=list)
    assumptions: list[str] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    sections: list[tuple[str, pd.DataFrame, str]] = field(default_factory=list)
    quality: pd.DataFrame | None = None
    charts: dict[str, Path] = field(default_factory=dict)
    meta: dict = field(default_factory=dict)


# --------------------------------------------------------------------------
# Форматирование
# --------------------------------------------------------------------------

def fmt_num(value, digits: int = 2) -> str:
    if value is None or (isinstance(value, float) and np.isnan(value)):
        return "—"
    if isinstance(value, (int, np.integer)):
        return f"{int(value):,}".replace(",", " ")
    try:
        number = float(value)
    except (TypeError, ValueError):
        return str(value)
    if abs(number) < 1 and number != 0:
        digits = max(digits, 4)
    text = f"{number:,.{digits}f}".replace(",", " ").replace(".", ",")
    return text


def _fmt_cell(value, column: str = "") -> str:
    if isinstance(value, (pd.Timestamp,)):
        return value.strftime("%d.%m.%Y")
    if isinstance(value, str):
        return value
    if column.endswith("_pct") or column in {"share_pct", "from_start_pct", "stage_pct"}:
        return "—" if pd.isna(value) else f"{float(value):.2f} %"
    if column in {"_z"}:
        return "—" if pd.isna(value) else f"{float(value):.2f}"
    return fmt_num(value)


def _table_for_output(df: pd.DataFrame) -> pd.DataFrame:
    """Привести таблицу к строкам для DOCX/PDF/PPTX (без потери чисел в XLSX)."""
    out = df.copy()
    for column in out.columns:
        out[column] = out[column].map(lambda value, c=column: _fmt_cell(value, c))
    return out


# --------------------------------------------------------------------------
# Шрифт для PDF
# --------------------------------------------------------------------------

def resolve_font() -> Path:
    """Найти TTF с кириллицей: assets → matplotlib → система."""
    candidates = [
        Path(__file__).resolve().parents[2] / "assets" / "fonts" / "DejaVuSans.ttf",
    ]
    try:
        import matplotlib
        candidates.append(Path(matplotlib.get_data_path()) / "fonts" / "ttf" / "DejaVuSans.ttf")
    except Exception:  # noqa: BLE001 - matplotlib может отсутствовать
        pass
    candidates += [
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        Path("C:/Windows/Fonts/DejaVuSans.ttf"),
        Path("/Library/Fonts/DejaVuSans.ttf"),
    ]
    for path in candidates:
        if path.exists():
            return path
    raise RuntimeError(
        "Не найден TTF с кириллицей. Установите DejaVu Sans или положите "
        "DejaVuSans.ttf в assets/fonts/."
    )


# --------------------------------------------------------------------------
# DOCX
# --------------------------------------------------------------------------

def _chart_title(name: str) -> str:
    return f"График: {name}"


def write_docx(report: ReportData, path: Path) -> Path:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt

    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Segoe UI"
    style.font.size = Pt(10)

    title = document.add_heading(report.title, level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if report.subtitle:
        subtitle = document.add_paragraph(report.subtitle)
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle.runs[0].font.size = Pt(11)
        subtitle.runs[0].font.color.rgb = None

    document.add_paragraph(
        f"Дата формирования: {pd.Timestamp.now().strftime('%d.%m.%Y %H:%M')} · "
        f"Версия скилла: {report.meta.get('version', '—')}"
    )

    if report.headline:
        document.add_heading("Ключевые показатели", level=1)
        table = document.add_table(rows=1, cols=3)
        table.style = "Light Grid Accent 1"
        headers = ("Показатель", "Значение", "Источник")
        for cell, text in zip(table.rows[0].cells, headers):
            cell.text = text
        for item in report.headline:
            row = table.add_row().cells
            row[0].text = str(item.get("label", ""))
            row[1].text = f"{item.get('value', '—')} {item.get('unit', '')}".strip()
            row[2].text = str(item.get("source", ""))

    if report.notes:
        document.add_heading("Источник данных", level=1)
        for item in report.notes:
            document.add_paragraph(item, style="List Bullet")

    if report.assumptions:
        document.add_heading("Допущения", level=1)
        for item in report.assumptions:
            document.add_paragraph(item, style="List Bullet")

    for name, frame, note in report.sections:
        document.add_heading(name, level=1)
        if note:
            document.add_paragraph(note)
        if frame is None or frame.empty:
            document.add_paragraph("Нет данных для этого раздела.")
            continue
        view = _table_for_output(frame)
        table = document.add_table(rows=1, cols=len(view.columns))
        try:
            table.style = "Light Grid Accent 1"
        except KeyError:
            table.style = "Table Grid"
        for cell, column in zip(table.rows[0].cells, view.columns):
            cell.text = str(column)
        for _, row in view.iterrows():
            cells = table.add_row().cells
            for cell, value in zip(cells, row.values):
                cell.text = str(value)

    if report.warnings:
        document.add_heading("Ограничения и предупреждения", level=1)
        for item in report.warnings:
            document.add_paragraph(item, style="List Bullet")

    if report.quality is not None and not report.quality.empty:
        document.add_heading("Контроль качества данных", level=1)
        view = _table_for_output(report.quality)
        table = document.add_table(rows=1, cols=len(view.columns))
        table.style = "Table Grid"
        for cell, column in zip(table.rows[0].cells, view.columns):
            cell.text = str(column)
        for _, row in view.iterrows():
            cells = table.add_row().cells
            for cell, value in zip(cells, row.values):
                cell.text = str(value)

    for name, chart in report.charts.items():
        if not chart or not Path(chart).exists():
            continue
        document.add_heading(_chart_title(name), level=1)
        document.add_picture(str(chart), width=Inches(PAGE_WIDTH_INCHES))

    # Номер страницы в колонтитуле
    footer = document.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run()
    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")
    run._r.append(field_begin)
    run._r.append(instruction)
    run._r.append(field_end)

    document.save(str(path))
    return path


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------

def write_pdf(report: ReportData, path: Path) -> Path:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Image, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    font_path = resolve_font()
    pdfmetrics.registerFont(TTFont("DejaVu", str(font_path)))

    styles = getSampleStyleSheet()
    for item in styles.byName.values():
        if hasattr(item, "fontName"):
            item.fontName = "DejaVu"
    styles.add(ParagraphStyle(name="RuTitle", parent=styles["Title"], fontName="DejaVu",
                              fontSize=20, leading=24, alignment=TA_CENTER, spaceAfter=10))
    styles.add(ParagraphStyle(name="RuHeading", parent=styles["Heading1"], fontName="DejaVu",
                              fontSize=14, leading=17, spaceBefore=12, spaceAfter=6))
    styles.add(ParagraphStyle(name="RuBody", parent=styles["BodyText"], fontName="DejaVu",
                              fontSize=9, leading=12))
    styles.add(ParagraphStyle(name="RuCell", parent=styles["BodyText"], fontName="DejaVu",
                              fontSize=7.5, leading=9.5))

    def page_number(canvas, doc):
        canvas.setFont("DejaVu", 8)
        canvas.setFillColor(colors.HexColor("#6b7280"))
        canvas.drawRightString(A4[0] - 40, 25, f"{doc.page}")

    document = SimpleDocTemplate(str(path), pagesize=A4,
                                 leftMargin=40, rightMargin=40, topMargin=40, bottomMargin=40)
    story: list = []
    story.append(Paragraph(escape(report.title), styles["RuTitle"]))
    if report.subtitle:
        story.append(Paragraph(escape(report.subtitle), styles["RuBody"]))
    story.append(Spacer(1, 8))
    story.append(Paragraph(
        f"Дата формирования: {pd.Timestamp.now().strftime('%d.%m.%Y %H:%M')} · "
        f"версия скилла {escape(str(report.meta.get('version', '—')))}",
        styles["RuBody"]))
    story.append(Spacer(1, 12))

    def make_table(frame: pd.DataFrame) -> Table:
        data = [list(frame.columns)] + frame.astype(str).values.tolist()
        table = Table(data, repeatRows=1, hAlign="LEFT")
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f3f4f6")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#374151")),
            ("FONTNAME", (0, 0), (-1, -1), "DejaVu"),
            ("FONTSIZE", (0, 0), (-1, -1), 7.5),
            ("LEADING", (0, 0), (-1, -1), 9.5),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#e5e7eb")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f9fafb")]),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        return table

    if report.headline:
        story.append(Paragraph("Ключевые показатели", styles["RuHeading"]))
        frame = pd.DataFrame(report.headline)[["label", "value", "unit", "source"]]
        frame.columns = ["Показатель", "Значение", "Ед.", "Источник"]
        story.append(make_table(frame))
        story.append(Spacer(1, 8))

    if report.notes:
        story.append(Paragraph("Источник данных", styles["RuHeading"]))
        for item in report.notes:
            story.append(Paragraph("• " + escape(item), styles["RuBody"]))

    if report.assumptions:
        story.append(Paragraph("Допущения", styles["RuHeading"]))
        for item in report.assumptions:
            story.append(Paragraph("• " + escape(item), styles["RuBody"]))

    for name, frame, note in report.sections:
        story.append(Paragraph(escape(name), styles["RuHeading"]))
        if note:
            story.append(Paragraph(escape(note), styles["RuBody"]))
        if frame is None or frame.empty:
            story.append(Paragraph("Нет данных для этого раздела.", styles["RuBody"]))
            continue
        story.append(make_table(_table_for_output(frame)))

    if report.warnings:
        story.append(Paragraph("Ограничения и предупреждения", styles["RuHeading"]))
        for item in report.warnings:
            story.append(Paragraph("• " + escape(item), styles["RuBody"]))

    if report.quality is not None and not report.quality.empty:
        story.append(Paragraph("Контроль качества данных", styles["RuHeading"]))
        story.append(make_table(_table_for_output(report.quality)))

    for name, chart in report.charts.items():
        if not chart or not Path(chart).exists():
            continue
        story.append(PageBreak())
        story.append(Paragraph(escape(_chart_title(name)), styles["RuHeading"]))
        story.append(Image(str(chart), width=6.9 * inch, height=3.4 * inch))

    document.build(story, onFirstPage=page_number, onLaterPages=page_number)
    return path


# --------------------------------------------------------------------------
# XLSX
# --------------------------------------------------------------------------

def write_xlsx(report: ReportData, path: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Отчёт"
    sheet["A1"] = report.title
    sheet["A1"].font = Font(bold=True, size=14)
    sheet["A2"] = f"Дата: {pd.Timestamp.now().strftime('%d.%m.%Y %H:%M')}"
    sheet["A2"].font = Font(size=9, color="6B7280")

    row = 4
    if report.headline:
        sheet.cell(row=row, column=1, value="Ключевые показатели").font = Font(bold=True, size=11)
        row += 1
        for column, header in enumerate(("Показатель", "Значение", "Ед.", "Источник"), start=1):
            cell = sheet.cell(row=row, column=column, value=header)
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", start_color="F3F4F6")
        row += 1
        for item in report.headline:
            sheet.cell(row=row, column=1, value=str(item.get("label", "")))
            value = item.get("value")
            cell = sheet.cell(row=row, column=2, value=value if isinstance(value, (int, float)) else str(value))
            if isinstance(value, (int, float)):
                cell.number_format = "# ##0.##"
            sheet.cell(row=row, column=3, value=str(item.get("unit", "")))
            sheet.cell(row=row, column=4, value=str(item.get("source", "")))
            row += 1
        row += 1

    for name, frame, note in report.sections:
        if frame is None or frame.empty:
            continue
        safe_name = re.sub(r"[\\/*?:\[\]]", "", name)[:31] or "Лист"
        ws = workbook.create_sheet(safe_name)
        ws.cell(row=1, column=1, value=name).font = Font(bold=True, size=12)
        start = 3
        if note:
            ws.cell(row=2, column=1, value=note).font = Font(size=9, italic=True, color="6B7280")
        for j, column in enumerate(frame.columns, start=1):
            cell = ws.cell(row=start, column=j, value=str(column))
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", start_color="F3F4F6")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for i, (_, source_row) in enumerate(frame.iterrows(), start=start + 1):
            for j, column in enumerate(frame.columns, start=1):
                value = source_row[column]
                if isinstance(value, (np.integer,)):
                    value = int(value)
                elif isinstance(value, (np.floating,)):
                    value = float(value)
                if isinstance(value, float) and np.isnan(value):
                    value = None
                cell = ws.cell(row=i, column=j, value=value)
                if isinstance(value, float):
                    cell.number_format = "# ##0.00"
                elif isinstance(value, int):
                    cell.number_format = "# ##0"
        for j, column in enumerate(frame.columns, start=1):
            width = max(10, min(38, len(str(column)) + 4))
            ws.column_dimensions[get_column_letter(j)].width = width
        ws.freeze_panes = ws.cell(row=start + 1, column=1)
        ws.auto_filter.ref = f"A{start}:{get_column_letter(len(frame.columns))}{start + len(frame)}"

    if report.quality is not None and not report.quality.empty:
        ws = workbook.create_sheet("Контроль качества")
        view = report.quality
        for j, column in enumerate(view.columns, start=1):
            cell = ws.cell(row=1, column=j, value=str(column))
            cell.font = Font(bold=True)
        for i, (_, source_row) in enumerate(view.iterrows(), start=2):
            for j, column in enumerate(view.columns, start=1):
                ws.cell(row=i, column=j, value=str(source_row[column]))

    workbook.save(str(path))
    return path


# --------------------------------------------------------------------------
# PPTX
# --------------------------------------------------------------------------

def write_pptx(report: ReportData, path: Path) -> Path:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)

    def add_title(slide, text: str, size: int = 28) -> None:
        box = slide.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(12.1), Inches(0.9))
        frame = box.text_frame
        frame.text = text
        paragraph = frame.paragraphs[0]
        paragraph.font.size = Pt(size)
        paragraph.font.bold = True
        paragraph.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)

    def add_footer(slide, index: int) -> None:
        box = slide.shapes.add_textbox(Inches(12.0), Inches(7.0), Inches(0.9), Inches(0.3))
        paragraph = box.text_frame.paragraphs[0]
        paragraph.text = str(index)
        paragraph.font.size = Pt(9)
        paragraph.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

    layout = presentation.slide_layouts[6]

    slide = presentation.slides.add_slide(layout)
    add_title(slide, report.title, size=34)
    if report.subtitle:
        box = slide.shapes.add_textbox(Inches(0.6), Inches(1.6), Inches(12.1), Inches(0.6))
        box.text_frame.text = report.subtitle
        box.text_frame.paragraphs[0].font.size = Pt(16)
    box = slide.shapes.add_textbox(Inches(0.6), Inches(6.6), Inches(9), Inches(0.4))
    box.text_frame.text = (f"{pd.Timestamp.now().strftime('%d.%m.%Y')} · "
                           f"скилл financial-analyst {report.meta.get('version', '')}")
    box.text_frame.paragraphs[0].font.size = Pt(11)
    add_footer(slide, 1)

    content_sections = [(name, frame) for name, frame, _ in report.sections if frame is not None and not frame.empty]
    charts = [(name, path) for name, path in report.charts.items() if path and Path(path).exists()]

    total = 1 + (1 if len(content_sections) + len(charts) > 4 else 0) + len(content_sections) + len(charts)
    if total > 4:
        slide = presentation.slides.add_slide(layout)
        add_title(slide, "Содержание", size=28)
        box = slide.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11), Inches(5))
        frame = box.text_frame
        frame.text = "\n".join([name for name, _ in content_sections]
                               + [_chart_title(name) for name, _ in charts])
        for paragraph in frame.paragraphs:
            paragraph.font.size = Pt(16)
        add_footer(slide, len(presentation.slides))

    for name, frame in content_sections:
        slide = presentation.slides.add_slide(layout)
        add_title(slide, name, size=24)
        rows = min(len(frame) + 1, 14)
        columns = min(len(frame.columns), 8)
        shape = slide.shapes.add_table(rows, columns, Inches(0.6), Inches(1.4),
                                       Inches(12.1), Inches(0.4 + 0.32 * rows))
        table = shape.table
        for j, column in enumerate(frame.columns[:columns]):
            table.cell(0, j).text = str(column)
            for paragraph in table.cell(0, j).text_frame.paragraphs:
                paragraph.font.size = Pt(10)
                paragraph.font.bold = True
        for i, (_, source_row) in enumerate(frame.head(rows - 1).iterrows(), start=1):
            for j, column in enumerate(frame.columns[:columns]):
                table.cell(i, j).text = _fmt_cell(source_row[column], column)
                for paragraph in table.cell(i, j).text_frame.paragraphs:
                    paragraph.font.size = Pt(9)
        add_footer(slide, len(presentation.slides))

    for name, chart in charts:
        slide = presentation.slides.add_slide(layout)
        add_title(slide, _chart_title(name), size=24)
        slide.shapes.add_picture(str(chart), Inches(0.8), Inches(1.4), width=Inches(11.7))
        add_footer(slide, len(presentation.slides))

    presentation.save(str(path))
    return path


# --------------------------------------------------------------------------
# RTF
# --------------------------------------------------------------------------

def _rtf_escape(text: str) -> str:
    """Экранировать всё вне ASCII по спецификации RTF (\\uN)."""
    out = []
    for char in str(text):
        code = ord(char)
        if char == "\n":
            out.append("\\par\n")
        elif code < 128:
            if char in "\\{}":
                out.append("\\" + char)
            else:
                out.append(char)
        else:
            out.append(f"\\u{code}?")
    return "".join(out)


def write_rtf(report: ReportData, path: Path) -> Path:
    lines: list[str] = ["{\\rtf1\\ansi\\deff0"]
    lines.append("{\\fonttbl{\\f0\\fnil\\fcharset0 Segoe UI;}{\\f1\\fmodern\\fcharset0 Courier New;}}")
    lines.append("\\f0\\fs22")

    def para(text: str, bold: bool = False, size: int = 22, font: int = 0) -> None:
        mark = "\\b " if bold else ""
        end = "\\b0" if bold else ""
        lines.append(f"\\f{font}\\fs{size} {mark}{_rtf_escape(text)}{end}\\par")

    para(report.title, bold=True, size=34)
    if report.subtitle:
        para(report.subtitle, size=22)
    para(f"Дата формирования: {pd.Timestamp.now().strftime('%d.%m.%Y %H:%M')}")
    lines.append("\\par")

    def table_text(frame: pd.DataFrame) -> str:
        return frame.to_string(index=False, max_colwidth=40)

    if report.headline:
        para("Ключевые показатели", bold=True, size=26)
        frame = pd.DataFrame(report.headline)
        frame = frame[["label", "value", "unit", "source"]]
        frame.columns = ["Показатель", "Значение", "Ед.", "Источник"]
        lines.append(f"\\f1\\fs16 {_rtf_escape(table_text(frame))}\\f0\\par")

    if report.notes:
        para("Источник данных", bold=True, size=26)
        for item in report.notes:
            para("• " + item)

    if report.assumptions:
        para("Допущения", bold=True, size=26)
        for item in report.assumptions:
            para("• " + item)

    for name, frame, note in report.sections:
        para(name, bold=True, size=26)
        if note:
            para(note)
        if frame is None or frame.empty:
            para("Нет данных для этого раздела.")
            continue
        lines.append(f"\\f1\\fs16 {_rtf_escape(table_text(_table_for_output(frame)))}\\f0\\par")

    if report.warnings:
        para("Ограничения и предупреждения", bold=True, size=26)
        for item in report.warnings:
            para("• " + item)

    if report.quality is not None and not report.quality.empty:
        para("Контроль качества данных", bold=True, size=26)
        lines.append(f"\\f1\\fs16 {_rtf_escape(table_text(report.quality))}\\f0\\par")

    lines.append("}")
    Path(path).write_text("\n".join(lines), encoding="ascii", errors="replace")
    return path


# --------------------------------------------------------------------------
# Проверка текста
# --------------------------------------------------------------------------

ALLOWED_LATIN = {
    "ltv", "cac", "cpl", "arpu", "arpa", "nim", "roa", "roe", "par", "npl", "dpd",
    "ecl", "kpi", "yoy", "qoq", "mom", "tat", "cagr", "roi", "romi", "cogs",
    "churn", "retention", "tofu", "mofu", "bofu", "abc", "iqr", "mape", "mae",
    "rmse", "arima", "ets", "csv", "excel", "json", "pdf", "docx", "xlsx",
    "pptx", "rtf", "html", "python", "sql", "github", "mit", "api", "id",
    "k-means", "z-score", "b2b", "b2c",
}


def lint_text(text: str) -> list[str]:
    """Проверить текст на смешение алфавитов, двойные пробелы и латиницу."""
    issues: list[str] = []
    for number, line in enumerate(text.splitlines(), start=1):
        if re.search(r"[ \t]{2,}\S", line) and not line.lstrip().startswith("|"):
            issues.append(f"строка {number}: двойной пробел")
        if line.rstrip() != line:
            issues.append(f"строка {number}: пробел в конце строки")
        for match in re.finditer(r"[A-Za-zА-Яа-яЁё]+", line):
            word = match.group()
            latin = any("LATIN" in unicodedata.name(ch, "") for ch in word)
            cyrillic = any("CYRILLIC" in unicodedata.name(ch, "") for ch in word)
            if latin and cyrillic:
                issues.append(f"строка {number}: смешение алфавитов в слове {word!r}")
        for match in re.finditer(r"\b[A-Za-z][A-Za-z-]{2,}\b", line):
            word = match.group()
            if word.lower() not in ALLOWED_LATIN:
                issues.append(f"строка {number}: латиница вне терминов — {word!r}")
    return issues


# --------------------------------------------------------------------------
# Сборка
# --------------------------------------------------------------------------

# --------------------------------------------------------------------------
# HTML (статический отчёт; интерактивный дашборд собирает visuals.py)
# --------------------------------------------------------------------------

def write_html(report: ReportData, path: Path) -> Path:
    import base64
    import html as html_lib

    def table_html(frame: pd.DataFrame) -> str:
        view = _table_for_output(frame)
        head = "".join(f"<th>{html_lib.escape(str(column))}</th>" for column in view.columns)
        body = "".join(
            "<tr>" + "".join(f"<td>{html_lib.escape(str(value))}</td>" for value in row) + "</tr>"
            for row in view.values
        )
        return f"<table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>"

    parts: list[str] = []
    if report.headline:
        parts.append("<h2>Ключевые показатели</h2>")
        parts.append(table_html(pd.DataFrame(report.headline)[["label", "value", "unit", "source"]]))
    if report.notes:
        parts.append("<h2>Источник данных</h2><ul>")
        parts.extend(f"<li>{html_lib.escape(item)}</li>" for item in report.notes)
        parts.append("</ul>")
    if report.assumptions:
        parts.append("<h2>Допущения</h2><ul>")
        parts.extend(f"<li>{html_lib.escape(item)}</li>" for item in report.assumptions)
        parts.append("</ul>")
    for name, frame, note in report.sections:
        parts.append(f"<h2>{html_lib.escape(name)}</h2>")
        if note:
            parts.append(f"<p class='note'>{html_lib.escape(note)}</p>")
        if frame is None or frame.empty:
            parts.append("<p>Нет данных для этого раздела.</p>")
        else:
            parts.append(table_html(frame))
    if report.warnings:
        parts.append("<h2>Ограничения и предупреждения</h2><ul>")
        parts.extend(f"<li>{html_lib.escape(item)}</li>" for item in report.warnings)
        parts.append("</ul>")
    if report.quality is not None and not report.quality.empty:
        parts.append("<h2>Контроль качества данных</h2>")
        parts.append(table_html(report.quality))
    for name, chart in report.charts.items():
        if chart and Path(chart).exists():
            encoded = base64.b64encode(Path(chart).read_bytes()).decode("ascii")
            parts.append(f"<h2>{html_lib.escape(_chart_title(name))}</h2>"
                         f"<img alt='{html_lib.escape(name)}' src='data:image/png;base64,{encoded}'>")

    document = f"""<!DOCTYPE html>
<html lang="ru"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{html_lib.escape(report.title)}</title>
<style>
 body {{ margin:0; padding:24px; background:#f8f9fa; color:#1f2937;
        font-family:"Segoe UI", Roboto, "Helvetica Neue", Arial, sans-serif; }}
 main {{ max-width:1200px; margin:0 auto; }}
 h1 {{ font-size:26px; margin:0 0 4px; }}
 h2 {{ font-size:17px; margin:26px 0 10px; color:#374151; }}
 .meta {{ color:#6b7280; font-size:13px; margin-bottom:18px; }}
 .note {{ color:#6b7280; font-size:13px; }}
 table {{ width:100%; border-collapse:collapse; background:#fff; font-size:13px; }}
 th {{ text-align:left; background:#f3f4f6; padding:9px; border-bottom:1px solid #e5e7eb; }}
 td {{ padding:8px 9px; border-bottom:1px solid #eef1f4; }}
 tr:hover td {{ background:#f9fafb; }}
 img {{ max-width:100%; height:auto; background:#fff; border:1px solid #e5e7eb; border-radius:8px; }}
 ul {{ background:#fff; border:1px solid #e5e7eb; border-radius:8px; padding:14px 14px 14px 30px; }}
 @media (max-width:640px) {{ body {{ padding:12px; }} table {{ font-size:12px; }} }}
</style></head>
<body><main>
<h1>{html_lib.escape(report.title)}</h1>
<div class="meta">{html_lib.escape(report.subtitle)} · сформировано {pd.Timestamp.now().strftime('%d.%m.%Y %H:%M')}</div>
{''.join(parts)}
</main></body></html>"""
    path.write_text(document, encoding="utf-8")
    return path


WRITERS = {
    "docx": write_docx,
    "pdf": write_pdf,
    "xlsx": write_xlsx,
    "pptx": write_pptx,
    "rtf": write_rtf,
}


def write_reports(report: ReportData, out_dir: Path, formats: list[str]) -> dict[str, Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    created: dict[str, Path] = {}
    unknown = sorted(set(formats) - set(WRITERS))
    if unknown:
        raise ValueError(f"Неизвестные форматы: {', '.join(unknown)}. Доступны: {', '.join(WRITERS)}")
    for name in formats:
        path = out_dir / f"financial_analyst_report.{name}"
        WRITERS[name](report, path)
        created[name] = path
    return created


def quality_frame(report: QualityReport) -> pd.DataFrame:
    return pd.DataFrame([
        {"Проверка": check.check_id, "Статус": check.status, "Детали": check.detail}
        for check in report.checks
    ])
