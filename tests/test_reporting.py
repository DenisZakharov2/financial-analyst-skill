import pandas as pd
import pytest
from finanalyst.reporting import ReportData, lint_text, quality_frame, write_html, write_reports
from finanalyst.visuals import chart_funnel, setup_matplotlib


@pytest.fixture(scope="module")
def report(tmp_path_factory):
    setup_matplotlib()
    out = tmp_path_factory.mktemp("charts")
    frame = pd.DataFrame({"stage": ["лид", "заявка"], "reached": [100, 40],
                          "from_start_pct": [100.0, 40.0], "stage_pct": [None, 40.0],
                          "lost": [None, 60]})
    chart = chart_funnel(frame, out / "funnel.png")
    return ReportData(
        title="Финансовый аналитический отчёт",
        subtitle="Тестовая выгрузка",
        headline=[{"label": "PAR 30", "value": "5,04", "raw": 5.04, "unit": "%",
                   "source": "portfolio.dpd"}],
        assumptions=["Этап считается максимальным достигнутым."],
        notes=["Строк в выгрузке: 120."],
        warnings=["В портфеле две валюты."],
        sections=[("Воронка продаж", frame, "Тестовое примечание")],
        quality=pd.DataFrame([{"Проверка": "pk", "Статус": "pass", "Детали": "OK"}]),
        charts={"Воронка": chart},
        meta={"data_dir": "fixtures", "version": "test"},
    )


def test_all_formats_created(report, tmp_path):
    created = write_reports(report, tmp_path, ["docx", "pdf", "xlsx", "pptx", "rtf"])
    assert set(created) == {"docx", "pdf", "xlsx", "pptx", "rtf"}
    for path in created.values():
        assert path.exists() and path.stat().st_size > 1000


def test_pdf_contains_cyrillic(report, tmp_path):
    pypdf = pytest.importorskip("pypdf")
    write_reports(report, tmp_path, ["pdf"])
    path = tmp_path / "financial_analyst_report.pdf"
    text = "".join(page.extract_text() for page in pypdf.PdfReader(path).pages)
    assert "Финансовый" in text
    assert "Допущения" in text
    assert "■" not in text


def test_rtf_is_ascii_and_escaped(report, tmp_path):
    write_reports(report, tmp_path, ["rtf"])
    data = (tmp_path / "financial_analyst_report.rtf").read_bytes()
    assert data.startswith(b"{\\rtf1")
    assert all(byte < 128 for byte in data)
    assert b"\\u1060?" in data


def test_xlsx_keeps_numbers_numeric(report, tmp_path):
    from openpyxl import load_workbook
    write_reports(report, tmp_path, ["xlsx"])
    workbook = load_workbook(tmp_path / "financial_analyst_report.xlsx")
    sheet = workbook["Воронка продаж"]
    values = [cell for row in sheet.iter_rows(min_row=4, max_row=4) for cell in row]
    assert any(cell.data_type == "n" for cell in values)


def test_xlsx_headline_value_is_numeric(report, tmp_path):
    from openpyxl import load_workbook
    write_reports(report, tmp_path, ["xlsx"])
    sheet = load_workbook(tmp_path / "financial_analyst_report.xlsx")["Отчёт"]
    cells = [cell for row in sheet.iter_rows(min_row=1, max_row=8, max_col=2) for cell in row]
    headline = [cell for cell in cells if cell.value == 5.04]
    assert headline and headline[0].data_type == "n"


def test_report_tables_use_russian_headers(report, tmp_path):
    from openpyxl import load_workbook
    write_reports(report, tmp_path, ["xlsx"])
    sheet = load_workbook(tmp_path / "financial_analyst_report.xlsx")["Воронка продаж"]
    headers = [cell.value for cell in sheet[3]]
    assert "Этап" in headers and "Дошедшие" in headers
    assert "stage" not in headers


def test_html_has_tables_and_images(report, tmp_path):
    path = write_html(report, tmp_path / "report.html")
    text = path.read_text(encoding="utf-8")
    assert "<table>" in text
    assert "data:image/png;base64" in text
    assert "@media" in text
    assert 'class="num"' in text


def test_quality_frame_has_russian_titles_and_codes():
    from finanalyst.core import QualityReport

    quality = QualityReport()
    quality.add("portfolio_pk_unique", True, "Дубли account_id: 0")
    quality.add("finance_required_columns", False, "Нет колонок product_code", warn_only=True)
    frame = quality_frame(quality)
    assert list(frame.columns) == ["Проверка", "Код", "Статус", "Детали"]
    assert frame.iloc[0]["Проверка"] == "Уникальность account_id"
    assert frame.iloc[0]["Статус"] == "пройдено"
    assert frame.iloc[1]["Проверка"] == "Обязательные колонки: finance"
    assert frame.iloc[1]["Код"] == "finance_required_columns"
    assert frame.iloc[1]["Статус"] == "предупреждение"


def test_lint_text_catches_problems():
    issues = lint_text("Кредитные карты cконцентрированы  в сегменте. Attracted средства.")
    joined = " ".join(issues)
    assert "смешение алфавитов" in joined
    assert "двойной пробел" in joined
    assert "Attracted" in joined


def test_lint_text_accepts_clean_text():
    assert lint_text("Портфель вырос на 5,04 %. LTV / CAC равен 3,2.") == []
